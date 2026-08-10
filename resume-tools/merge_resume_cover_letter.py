import argparse
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pypdf import PdfReader, PdfWriter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


NAVY = "13283F"
TEAL = "159A9C"
GRAY = "5C6670"
LIGHT_GRAY = "D9E0E6"
FONT_REGULAR = Path(r"C:\Windows\Fonts\malgun.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\malgunbd.ttf")
EXCLUDED_HEADING = "지원 분야별 활용 방향"


def set_run_style(run, size, bold=False, color=NAVY):
    """DOCX 런에 이력서와 동일한 한글 글꼴·크기·색상을 적용합니다."""
    run.font.name = "맑은 고딕"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "맑은 고딕")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "맑은 고딕")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_bottom_border(paragraph, color=LIGHT_GRAY, size="8", space="5"):
    """문단 아래에 절제된 구분선을 추가해 자기소개서 제목 계층을 구분합니다."""
    paragraph_properties = paragraph._p.get_or_add_pPr()
    borders = paragraph_properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        paragraph_properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def parse_cover_letter(markdown_path):
    """자기소개서 Markdown에서 1~6번 문항만 추출하고 안내용 후반부를 제외합니다."""
    text = markdown_path.read_text(encoding="utf-8")
    text = text.split(f"## {EXCLUDED_HEADING}", maxsplit=1)[0]
    section_pattern = re.compile(r"^##\s+(\d+)\.\s+(.+?)\s*$", re.MULTILINE)
    matches = list(section_pattern.finditer(text))
    sections = []
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        body = text[start:end].strip()
        paragraphs = [item.strip() for item in re.split(r"\n\s*\n", body) if item.strip() and item.strip() != "---"]
        sections.append((int(match.group(1)), match.group(2).strip(), paragraphs))
    if len(sections) != 6:
        raise ValueError(f"자기소개서 1~6번 문항을 모두 찾지 못했습니다: {len(sections)}개")
    return sections


def add_cover_letter_pages(resume_path, sections, output_path):
    """기존 이력서 뒤에 문항을 연속 배치한 통합 DOCX를 생성합니다."""
    document = Document(resume_path)
    document.add_page_break()
    for index, (number, title, body_paragraphs) in enumerate(sections):
        if index == 0:
            title_paragraph = document.add_paragraph()
            title_paragraph.paragraph_format.space_after = Pt(2)
            title_run = title_paragraph.add_run("자기소개서")
            set_run_style(title_run, 22, bold=True)

            identity_paragraph = document.add_paragraph()
            identity_paragraph.paragraph_format.space_after = Pt(22)
            identity_run = identity_paragraph.add_run("김종록  |  BACKEND & AI APPLICATION DEVELOPER")
            set_run_style(identity_run, 9, bold=True, color=TEAL)

        heading = document.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.page_break_before = number == 5
        heading.paragraph_format.space_before = Pt(18 if index else 0)
        heading.paragraph_format.space_after = Pt(15)
        number_run = heading.add_run(f"{number:02d}  ")
        set_run_style(number_run, 14, bold=True, color=TEAL)
        heading_run = heading.add_run(title)
        set_run_style(heading_run, 14, bold=True)
        add_bottom_border(heading)

        for paragraph_text in body_paragraphs:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.line_spacing = 1.4
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.widow_control = True
            body_run = paragraph.add_run(paragraph_text)
            set_run_style(body_run, 9.2)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def register_pdf_fonts():
    """통합 PDF의 자기소개서 페이지에 사용할 맑은 고딕 글꼴을 등록합니다."""
    pdfmetrics.registerFont(TTFont("Malgun", str(FONT_REGULAR)))
    pdfmetrics.registerFont(TTFont("MalgunBold", str(FONT_BOLD)))


def wrap_text(text, font_name, font_size, max_width):
    """ReportLab 글꼴 폭을 기준으로 한국어 본문을 페이지 폭 안에서 줄바꿈합니다."""
    words = text.split()
    lines = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if current and pdfmetrics.stringWidth(candidate, font_name, font_size) > max_width:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def draw_continuous_cover_letter(page_width, page_height, sections, first_page_number):
    """자기소개서 문항을 불필요한 강제 페이지 나눔 없이 연속 PDF로 생성합니다."""
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(page_width, page_height))
    left = 44.3
    right = page_width - 44.3
    bottom = 55
    page_number = first_page_number

    def start_page(first_page=False):
        """새 자기소개서 페이지의 머리말과 선택적 표지를 그리고 본문 시작 좌표를 반환합니다."""
        y = page_height - 42
        pdf.setFont("MalgunBold", 7.8)
        pdf.setFillColor(f"#{GRAY}")
        pdf.drawString(left, y, "KIM JONGROK · COVER LETTER")
        y -= 31
        if not first_page:
            return y
        pdf.setFont("MalgunBold", 22)
        pdf.setFillColor(f"#{NAVY}")
        pdf.drawString(left, y, "자기소개서")
        y -= 23
        pdf.setFont("MalgunBold", 8.5)
        pdf.setFillColor(f"#{TEAL}")
        pdf.drawString(left, y, "BACKEND & AI APPLICATION DEVELOPER")
        y -= 38
        return y

    def finish_page(current_page_number):
        """현재 페이지에 쪽 번호를 추가하고 다음 PDF 페이지로 전환합니다."""
        pdf.setFont("Malgun", 7.5)
        pdf.setFillColor(f"#{GRAY}")
        pdf.drawRightString(right, 25, str(current_page_number))
        pdf.showPage()

    font_size = 9.2
    line_height = 15.4
    max_width = right - left
    y = start_page(first_page=True)
    for section_index, (number, title, paragraphs) in enumerate(sections):
        wrapped_paragraphs = [wrap_text(text, "Malgun", font_size, max_width) for text in paragraphs]
        heading_height = 58 if section_index else 40
        if number == 5 and y < page_height - 73:
            finish_page(page_number)
            page_number += 1
            y = start_page()
        elif y - heading_height < bottom:
            finish_page(page_number)
            page_number += 1
            y = start_page()
        elif section_index:
            y -= 18

        pdf.setFont("MalgunBold", 13.5)
        pdf.setFillColor(f"#{TEAL}")
        pdf.drawString(left, y, f"{number:02d}")
        pdf.setFillColor(f"#{NAVY}")
        pdf.drawString(left + 34, y, title)
        y -= 13
        pdf.setStrokeColor(f"#{LIGHT_GRAY}")
        pdf.setLineWidth(0.8)
        pdf.line(left, y, right, y)
        y -= 24

        for lines in wrapped_paragraphs:
            for line in lines:
                if y < bottom:
                    finish_page(page_number)
                    page_number += 1
                    y = start_page()
                pdf.setFont("Malgun", font_size)
                pdf.setFillColor(f"#{NAVY}")
                pdf.drawString(left, y, line)
                y -= line_height
            y -= 6
        y -= 6

    finish_page(page_number)
    pdf.save()
    buffer.seek(0)
    return buffer


def build_combined_pdf(resume_pdf_path, sections, output_path):
    """기존 이력서 PDF 뒤에 연속 배치한 자기소개서를 병합합니다."""
    register_pdf_fonts()
    resume_reader = PdfReader(resume_pdf_path)
    writer = PdfWriter()
    for page in resume_reader.pages:
        writer.add_page(page)
    page_width = float(resume_reader.pages[0].mediabox.width)
    page_height = float(resume_reader.pages[0].mediabox.height)
    cover_letter_buffer = draw_continuous_cover_letter(
        page_width,
        page_height,
        sections,
        first_page_number=len(resume_reader.pages) + 1,
    )
    for page in PdfReader(cover_letter_buffer).pages:
        writer.add_page(page)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("wb") as output_file:
        writer.write(output_file)


def main():
    """명령행 인자를 읽어 이력서·자기소개서 통합 DOCX와 PDF를 함께 생성합니다."""
    parser = argparse.ArgumentParser()
    parser.add_argument("--resume-docx", type=Path, required=True)
    parser.add_argument("--resume-pdf", type=Path, required=True)
    parser.add_argument("--cover-letter", type=Path, required=True)
    parser.add_argument("--output-docx", type=Path, required=True)
    parser.add_argument("--output-pdf", type=Path, required=True)
    args = parser.parse_args()

    sections = parse_cover_letter(args.cover_letter)
    add_cover_letter_pages(args.resume_docx, sections, args.output_docx)
    build_combined_pdf(args.resume_pdf, sections, args.output_pdf)
    print(args.output_docx)
    print(args.output_pdf)


if __name__ == "__main__":
    main()
