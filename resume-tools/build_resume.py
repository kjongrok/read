from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from docx.text.run import Run


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output"
OUTPUT_PATH = OUTPUT_DIR / "김종록_백엔드_AI응용개발자_이력서.docx"

NAVY = "13283F"
TEAL = "159A9C"
LIGHT_TEAL = "EAF6F5"
GRAY = "5C6670"
LIGHT_GRAY = "D9E0E6"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    """표 셀의 배경색을 지정합니다.

    Args:
        cell: 배경색을 적용할 python-docx 셀입니다.
        fill: 16진수 RGB 색상 문자열입니다.

    Returns:
        반환값은 없습니다.

    Side effects:
        셀 XML에 배경색 속성을 추가합니다.
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges):
    """표 셀의 테두리 속성을 방향별로 설정합니다.

    Args:
        cell: 테두리를 적용할 python-docx 셀입니다.
        **edges: top, bottom, left, right 방향별 테두리 설정입니다.

    Returns:
        반환값은 없습니다.

    Side effects:
        셀 XML의 테두리 속성을 변경합니다.
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge_name, edge_data in edges.items():
        edge = tc_borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            tc_borders.append(edge)
        for key, value in edge_data.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row):
    """표의 첫 행이 다음 페이지에서도 제목 행으로 반복되도록 설정합니다.

    Args:
        row: 반복 대상으로 지정할 python-docx 행입니다.

    Returns:
        반환값은 없습니다.

    Side effects:
        행 XML에 반복 제목 속성을 추가합니다.
    """
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    """표 셀 안쪽 여백을 트윕 단위로 설정합니다.

    Args:
        cell: 여백을 적용할 python-docx 셀입니다.
        top: 위쪽 여백입니다.
        start: 시작 방향 여백입니다.
        bottom: 아래쪽 여백입니다.
        end: 끝 방향 여백입니다.

    Returns:
        반환값은 없습니다.

    Side effects:
        셀 XML의 여백 속성을 변경합니다.
    """
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def style_run(run, size=9.2, bold=False, color=NAVY, font="맑은 고딕"):
    """문서 전반에서 사용하는 글꼴과 색상을 런에 적용합니다.

    Args:
        run: 서식을 적용할 python-docx 런입니다.
        size: 글자 크기(pt)입니다.
        bold: 굵게 표시할지 여부입니다.
        color: 16진수 RGB 색상 문자열입니다.
        font: 적용할 글꼴 이름입니다.

    Returns:
        서식이 적용된 런을 반환합니다.

    Side effects:
        런의 글꼴, 크기, 굵기, 색상을 변경합니다.
    """
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_hyperlink(paragraph, text, url, size=8.7, bold=False, color=NAVY):
    """문단에 외부 URL이 연결된 하이퍼링크 런을 추가합니다.

    Args:
        paragraph: 링크를 추가할 python-docx 문단입니다.
        text: 문서에 표시할 링크 문자열입니다.
        url: 클릭 시 이동할 외부 URL입니다.
        size: 링크 글자 크기(pt)입니다.
        bold: 링크를 굵게 표시할지 여부입니다.
        color: 링크에 적용할 16진수 RGB 색상입니다.

    Returns:
        생성한 python-docx 런을 반환합니다.

    Side effects:
        문서 관계 파일에 외부 링크를 등록하고 문단 XML에 하이퍼링크를 추가합니다.
    """
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run_element = OxmlElement("w:r")
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)
    run = Run(run_element, paragraph)
    return style_run(run, size=size, bold=bold, color=color)


def add_portfolio_link(document, label, url, size=8.6, space_after=0.8):
    """굵은 항목명과 클릭 가능한 URL로 구성된 포트폴리오 문단을 추가합니다.

    Args:
        document: 내용을 추가할 python-docx 문서입니다.
        label: 굵게 표시할 링크 항목명입니다.
        url: 표시하고 연결할 외부 URL입니다.
        size: 문단 글자 크기(pt)입니다.
        space_after: 문단 아래 여백(pt)입니다.

    Returns:
        생성한 문단을 반환합니다.

    Side effects:
        문서에 새 문단과 외부 하이퍼링크 관계를 추가합니다.
    """
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.08
    style_run(paragraph.add_run(label), size=size, bold=True)
    add_hyperlink(paragraph, f"  {url}", url, size=size)
    return paragraph


def configure_document(document):
    """A4 이력서에 맞는 여백과 기본 스타일을 설정합니다.

    Args:
        document: 설정할 python-docx 문서입니다.

    Returns:
        반환값은 없습니다.

    Side effects:
        문서 섹션과 Normal 스타일을 변경합니다.
    """
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)

    normal = document.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(2.4)
    normal.paragraph_format.line_spacing = 1.14


def add_header(document):
    """이름, 목표 직무, 연락처가 포함된 첫 페이지 헤더를 작성합니다.

    Args:
        document: 내용을 추가할 python-docx 문서입니다.

    Returns:
        반환값은 없습니다.

    Side effects:
        문서 첫 부분에 이름과 연락처 문단을 추가합니다.
    """
    name_paragraph = document.add_paragraph()
    name_paragraph.paragraph_format.space_after = Pt(0)
    style_run(name_paragraph.add_run("김종록"), size=25, bold=True, color=NAVY)

    role_paragraph = document.add_paragraph()
    role_paragraph.paragraph_format.space_before = Pt(0)
    role_paragraph.paragraph_format.space_after = Pt(5)
    style_run(role_paragraph.add_run("BACKEND & AI APPLICATION DEVELOPER"), size=10.5, bold=True, color=TEAL)

    contact = document.add_paragraph()
    contact.paragraph_format.space_after = Pt(1)
    contact.paragraph_format.line_spacing = 1.0
    add_hyperlink(contact, "010-7742-1623", "tel:+821077421623", size=8.7, color=GRAY)
    style_run(contact.add_run("  |  "), size=8.7, color=GRAY)
    add_hyperlink(contact, "xhxhahs2@gmail.com", "mailto:xhxhahs2@gmail.com", size=8.7, color=GRAY)
    style_run(contact.add_run("  |  경기 하남시"), size=8.7, color=GRAY)

    rule = document.add_paragraph()
    rule.paragraph_format.space_after = Pt(6)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "16")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), TEAL)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_section_title(document, title):
    """구분선이 포함된 이력서 섹션 제목을 추가합니다.

    Args:
        document: 내용을 추가할 python-docx 문서입니다.
        title: 표시할 섹션 제목입니다.

    Returns:
        생성한 제목 문단을 반환합니다.

    Side effects:
        문서에 새 제목 문단을 추가합니다.
    """
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3.5)
    paragraph.paragraph_format.keep_with_next = True
    style_run(paragraph.add_run(title), size=11.3, bold=True, color=NAVY)
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "7")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), LIGHT_GRAY)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    return paragraph


def add_body_paragraph(document, text, bold_prefix=None, color=NAVY, size=9.2, space_after=2.5):
    """일반 본문 문단을 추가하고 선택적으로 앞부분을 강조합니다.

    Args:
        document: 내용을 추가할 python-docx 문서입니다.
        text: 문단 전체 문자열입니다.
        bold_prefix: 굵게 표시할 시작 문자열입니다.
        color: 본문 글자 색상입니다.
        size: 글자 크기(pt)입니다.
        space_after: 문단 뒤 간격(pt)입니다.

    Returns:
        생성한 본문 문단을 반환합니다.

    Side effects:
        문서에 새 본문 문단을 추가합니다.
    """
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.13
    if bold_prefix and text.startswith(bold_prefix):
        style_run(paragraph.add_run(bold_prefix), size=size, bold=True, color=color)
        style_run(paragraph.add_run(text[len(bold_prefix):]), size=size, color=color)
    else:
        style_run(paragraph.add_run(text), size=size, color=color)
    return paragraph


def add_bullet(document, text, level=0, size=8.9, color=NAVY, compact=True):
    """성과 중심의 글머리표 문단을 추가합니다.

    Args:
        document: 내용을 추가할 python-docx 문서입니다.
        text: 표시할 글머리표 내용입니다.
        level: 들여쓰기 단계입니다.
        size: 글자 크기(pt)입니다.
        color: 글자 색상입니다.
        compact: 문단 간격을 좁게 유지할지 여부입니다.

    Returns:
        생성한 글머리표 문단을 반환합니다.

    Side effects:
        문서에 새 글머리표 문단을 추가합니다.
    """
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.45 + 0.35 * level)
    paragraph.paragraph_format.first_line_indent = Cm(-0.22)
    paragraph.paragraph_format.space_after = Pt(1.2 if compact else 2.6)
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.paragraph_format.keep_together = True
    style_run(paragraph.add_run(text), size=size, color=color)
    return paragraph


def add_skill_table(document):
    """기술을 역할별로 분류한 간결한 표를 추가합니다.

    Args:
        document: 내용을 추가할 python-docx 문서입니다.

    Returns:
        생성한 기술 표를 반환합니다.

    Side effects:
        문서에 2열 기술 분류 표를 추가합니다.
    """
    rows = [
        ("Backend", "Java 17 · Spring Boot 3.5 · Spring Security · Spring Data JPA · MyBatis · SSE"),
        ("AI · ML/DL", "Python 3.12 · FastAPI · scikit-learn · TensorFlow/Keras · PyTorch · YOLOv8 · RAG · BGE-M3 · Qwen/Ollama"),
        ("Frontend", "React 18 · Vite 5 · React Router · Axios · Streamlit"),
        ("Data · Infra", "PostgreSQL · pgvector · Redis · MariaDB · Supabase · Docker(실행·검증)"),
        ("Test · Collaboration", "JUnit 5 · pytest · Bruno · Git/GitHub · GitHub Actions · Jira · Slack"),
    ]
    table = document.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(3.0)
    table.columns[1].width = Cm(14.4)
    for index, (label, values) in enumerate(rows):
        left, right = table.rows[index].cells
        left.width = Cm(3.0)
        right.width = Cm(14.4)
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(left, LIGHT_TEAL)
        for cell in (left, right):
            set_cell_margins(cell, top=60, bottom=60, start=100, end=100)
            set_cell_border(
                cell,
                bottom={"val": "single", "sz": "4", "color": LIGHT_GRAY},
            )
        left_p = left.paragraphs[0]
        left_p.paragraph_format.space_after = Pt(0)
        style_run(left_p.add_run(label), size=8.6, bold=True, color=TEAL)
        right_p = right.paragraphs[0]
        right_p.paragraph_format.space_after = Pt(0)
        style_run(right_p.add_run(values), size=8.5, color=NAVY)
    return table


def add_entry_heading(document, title, meta):
    """경력이나 프로젝트의 제목과 기간·역할 정보를 한 줄로 추가합니다.

    Args:
        document: 내용을 추가할 python-docx 문서입니다.
        title: 항목 제목입니다.
        meta: 기간과 역할 등 보조 정보입니다.

    Returns:
        생성한 제목 문단을 반환합니다.

    Side effects:
        문서에 새 항목 제목 문단을 추가합니다.
    """
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3.5)
    paragraph.paragraph_format.space_after = Pt(1.4)
    paragraph.paragraph_format.keep_with_next = True
    style_run(paragraph.add_run(title), size=10.0, bold=True, color=NAVY)
    style_run(paragraph.add_run(f"  |  {meta}"), size=8.4, color=GRAY)
    return paragraph


def add_page_number(section):
    """문서 바닥글 오른쪽에 현재 페이지 번호를 추가합니다.

    Args:
        section: 바닥글을 설정할 문서 섹션입니다.

    Returns:
        반환값은 없습니다.

    Side effects:
        섹션 바닥글에 PAGE 필드를 추가합니다.
    """
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    style_run(run, size=7.5, color=GRAY)


def build_resume():
    """검토된 경력과 프로젝트 사실을 바탕으로 2페이지 이력서를 생성합니다.

    Args:
        인자는 없습니다.

    Returns:
        생성된 DOCX 파일 경로를 반환합니다.

    Side effects:
        read/output 폴더에 기존 파일을 덮어써 DOCX 파일을 저장합니다.

    Raises:
        출력 폴더 생성이나 DOCX 저장에 실패하면 파일 시스템 예외가 발생합니다.
    """
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    add_header(document)

    add_section_title(document, "PROFILE")
    add_body_paragraph(
        document,
        "병원 전산 운영 3년 11개월 동안 여러 진료·행정부서의 장애에 대응하고 사용자 요청을 처리하며, 문제를 단순 증상이 아닌 업무 흐름과 시스템 구조의 관점에서 파악하는 경험을 쌓았습니다.",
        size=9.1,
        space_after=2.2,
    )
    add_body_paragraph(
        document,
        "이후 Spring Boot 기반 Backend와 React 화면을 구현하고, 머신러닝·딥러닝 모델 개발과 LLM 에이전트·RAG 챗봇 구현을 경험했습니다. 5인 팀 PL로 참여한 BidMatch에서는 FastAPI AI 기능을 Backend 규칙, 응답 검증과 대체 처리로 보완하고 AI 자동화 테스트 242건을 통과했습니다.",
        size=9.1,
        space_after=2.2,
    )
    add_body_paragraph(
        document,
        "운영 환경에 대한 이해를 바탕으로 AI 기능을 안정적인 Backend 서비스에 연결하고, 예외 상황에서도 사용자가 신뢰할 수 있는 결과를 제공하는 개발자로 성장하고 있습니다.",
        size=9.1,
        space_after=2.0,
    )

    links = document.add_paragraph()
    links.paragraph_format.space_after = Pt(4)
    links.paragraph_format.line_spacing = 1.0
    add_hyperlink(links, "Portfolio", "https://kjongrok.github.io/", size=8.7, bold=True, color=TEAL)
    style_run(links.add_run("  |  "), size=8.7, color=GRAY)
    add_hyperlink(links, "GitHub", "https://github.com/kjongrok", size=8.7, bold=True, color=TEAL)
    style_run(links.add_run("  |  "), size=8.7, color=GRAY)
    add_hyperlink(links, "BidMatch Project", "https://github.com/aiHuman1Team", size=8.7, bold=True, color=TEAL)
    style_run(links.add_run("  |  "), size=8.7, color=GRAY)
    add_hyperlink(links, "BidMatch Service", "http://minsworkspace.ddns.net", size=8.7, bold=True, color=TEAL)

    add_section_title(document, "CORE COMPETENCIES")
    add_bullet(document, "Spring Boot 기반 인증·회원·기업정보·고객센터·알림 API와 React 사용자 흐름 구현")
    add_bullet(document, "정형 데이터 ML, 영상·시계열 DL 모델 개발과 LLM 에이전트·RAG 챗봇 구현")
    add_bullet(document, "BGE-M3 검색·Qwen 분류와 Spring 규칙 판정을 분리하고 검증·fallback 적용")
    add_bullet(document, "5인 팀 프로젝트 리더(PL)로 API 계약, 기능 범위, 테스트와 시연 자료를 조율")

    add_section_title(document, "TECHNICAL SKILLS")
    add_skill_table(document)

    add_section_title(document, "WORK EXPERIENCE")
    add_entry_heading(document, "㈜시스게이트 · 한림대학교의료원(강동) 전산 운영", "2019.08 – 2023.06 · OA 운영 및 사용자 지원")
    add_body_paragraph(document, "담당 업무", size=8.5, color=TEAL, space_after=0.5)
    add_bullet(document, "병원 내 여러 진료·행정부서의 OA 요구사항 수집, 시스템 구성 검토와 도입 지원", size=8.6)
    add_bullet(document, "장애·사용자 요청 대응, 신규 시스템 교육 자료 제작과 사용자 교육 진행", size=8.6)
    add_body_paragraph(document, "경험 및 직무 연결점", size=8.5, color=TEAL, space_after=0.5)
    add_bullet(document, "반복 문의를 운영 매뉴얼·정책·Q&A 대응 체계로 정리하고 원격 OA 전환 지원", size=8.6)
    add_bullet(document, "사용자 업무와 영향 범위를 고려해 원인을 좁히는 경험을 개발의 예외 처리·운영 관점으로 확장", size=8.6)

    add_section_title(document, "EDUCATION")
    add_entry_heading(document, "휴먼AI교육센터 · 심화_인공지능(AI) 서비스 기반 웹 개발자 심화 프로젝트", "2026.06.11 – 2026.08.11")
    add_entry_heading(document, "KG IT BANK · 핀테크 서비스를 위한 풀스택 개발자 양성 과정", "2024.08.05 – 2025.02.13")
    add_entry_heading(document, "여주대학교 · 컴퓨터정보과", "2016.03 – 2020.02 · 졸업")

    document.add_page_break()

    add_section_title(document, "AI EDUCATION HIGHLIGHTS")
    add_bullet(document, "scikit-learn 기반 데이터 전처리·모델 비교와 불균형 데이터 평가")
    add_bullet(document, "YOLOv8·LSTM 기반 객체 탐지, 이상탐지와 시계열 예측 모델 개발")
    add_bullet(document, "LLM Function Calling·RAG 챗봇과 응답 검증·fallback 구현")
    add_bullet(document, "Spring Boot·FastAPI·React 기반 Backend·AI·사용자 화면 연동")

    add_section_title(document, "SELECTED PROJECTS")
    add_entry_heading(document, "BidMatch · AI 기반 공공입찰 맞춤 추천·자가 자격 진단 서비스", "2026.07.09 – 2026.08.07 · 5인 팀 PL")
    add_body_paragraph(
        document,
        "나라장터 공고와 기업정보를 연결해 맞춤 공고 추천, 근거 기반 자가 자격 진단, 알림과 고객센터를 제공하는 서비스",
        size=8.9,
        color=GRAY,
        space_after=2.0,
    )
    add_bullet(document, "Java 17·Spring Boot로 일반/기업회원 인증, 기업정보·면허·실적, 고객센터와 SSE 알림 API 구현")
    add_bullet(document, "React 사용자 화면을 연동하고 Bruno 요청으로 Backend·AI API 계약과 예외 응답 검증")
    add_bullet(document, "FAQ RAG 챗봇에 역할별 검색, 낮은 유사도 차단, 다중 의도, LLM fallback과 민감정보 마스킹 적용")
    add_bullet(document, "BGE-M3 근거 검색과 Qwen 유형 분류 결과를 Spring 규칙 판정으로 재검증하는 자가 자격 진단 구조 설계")
    add_bullet(document, "병렬 개발 중 DB 마이그레이션(Flyway) 버전·체크섬·스키마 충돌을 기존 이력을 보존한 후속 마이그레이션으로 해결")
    add_bullet(document, "5개 저장소 118개 비병합 커밋 기여, AI 자동화 테스트 242건, Backend 테스트, 사용자·관리자 Frontend 빌드 통과")
    add_body_paragraph(document, "기술: Spring Boot · FastAPI · React · PostgreSQL · Redis · BGE-M3 · Qwen/Ollama · RAG · PyMuPDF · Docker(실행·검증)", size=8.2, color=TEAL, space_after=2.0)

    add_entry_heading(document, "지능형 교통 관제 시스템", "2026.07.02 – 2026.07.07 · 개인")
    add_bullet(document, "YOLOv8 차량 탐지 결과와 Supabase 로그 조회를 LLM Function Calling으로 연결")
    add_bullet(document, "날짜 환각과 도구 호출 오류를 시간 파서·DB 결과 검증·대체 모델 경로로 완화하고 Word 관제 일지 생성")
    add_body_paragraph(document, "기술: Python · YOLOv8 · Supabase · LLM Function Calling · Streamlit · python-docx", size=8.0, color=TEAL, space_after=1.0)

    add_entry_heading(document, "CCTV 교통량 이상탐지·예측", "2026.06.25 – 2026.07.01 · 개인")
    add_bullet(document, "YOLOv8 탐지, LSTM Autoencoder 이상탐지, ITS CCTV, Supabase 로그를 Streamlit 화면으로 통합")
    add_bullet(document, "입력 차원 불일치와 영상 지연을 reshape, 프레임 스킵, 표시 영상 리사이징으로 해결")
    add_body_paragraph(document, "기술: Python · YOLOv8 · TensorFlow/Keras · LSTM Autoencoder · Supabase · Streamlit", size=8.0, color=TEAL, space_after=1.0)

    add_entry_heading(document, "항공편 지연 예측", "2026.06.18 – 2026.06.24 · 개인")
    add_bullet(document, "약 25만 건 데이터를 전처리하고 XGBoost 등 모델을 비교해 불균형 데이터의 지연 Recall 중심으로 평가")
    add_bullet(document, "전처리·모델을 scikit-learn Pipeline으로 패키징해 Streamlit에 연결, 프로젝트 보고서 기준 Recall 65%·ROC-AUC 0.785 기록")
    add_body_paragraph(document, "기술: Python · pandas · scikit-learn · XGBoost · Pipeline · Streamlit", size=8.0, color=TEAL, space_after=1.0)

    add_entry_heading(document, "맞춤 공고 자동 수집·메일 알림 프로토타입", "2026.06.11 – 2026.06.17 · 팀 PL")
    add_bullet(document, "Flask·React로 OAuth/JWT 인증, Gmail 알림, Gemini 공고 요약, APScheduler 배치를 연결")
    add_bullet(document, "UTC/KST 발송 시간 차이를 Asia/Seoul 기준으로 보정하고 최종 Spring·FastAPI 프로젝트로 구조 확장")
    add_body_paragraph(document, "기술: Flask · React · OAuth/JWT · Gmail SMTP · Gemini · APScheduler", size=8.0, color=TEAL, space_after=1.0)

    add_entry_heading(document, "Spike · Spring Boot 기반 금융 서비스", "2025.01 – 2025.02 · 팀 프로젝트")
    add_bullet(document, "Spring Security와 BCrypt를 적용한 회원가입·로그인, 역할별 접근 제어와 마이페이지 구현")
    add_bullet(document, "관리자 사용자 CRUD·페이징과 보이스피싱 의심 계좌 신고·검토·상태 변경 흐름 구현")
    add_body_paragraph(document, "기술: Java 17 · Spring Boot 2.7 · JSP · JPA · MyBatis · Oracle DB", size=8.2, color=TEAL, space_after=2.0)

    add_page_number(document.sections[0])
    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_resume())
